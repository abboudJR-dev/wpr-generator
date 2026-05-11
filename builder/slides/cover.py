"""Slides 1-3: Cover, Project Overview, Project Information Summary."""
from __future__ import annotations

from pptx.util import Inches, Pt

from builder.design import (
    ASSETS, C, Run, SLIDE_H, SLIDE_W,
    add_blank_slide, add_footer, add_header, add_image, add_oval, add_rect,
    add_text_box, fill_cell, hex_to_rgb, kpi_card_inline,
    section_label, set_cell_border, set_slide_background, style_cell_text,
)
from builder.state import WPRState


def _add_kpi(slide, x, y, w, h, value, label, fill, value_size=22):
    """Centered value+label KPI (used on slides 3 & 9)."""
    add_rect(slide, x, y, w, h, fill)
    add_text_box(
        slide, x, y + 0.10, w, h * 0.55,
        [Run(text=value, size=value_size, bold=True, color=C.white)],
        align="center", valign="middle", margin=0,
    )
    add_text_box(
        slide, x, y + h - 0.30, w, 0.25,
        [Run(text=label.upper(), size=8.5, color=C.white, char_spacing=2.5)],
        align="center", valign="middle", margin=0,
    )


# ---------- SLIDE 1 — COVER ---------------------------------------------------

def build_slide_cover(prs, state: WPRState):
    slide = add_blank_slide(prs)
    set_slide_background(slide, C.navy_deep)

    # Building rendering — top half hero
    add_image(slide, ASSETS["building"], 0, 0, SLIDE_W, 4.0)

    # Teal accent bar under hero
    add_rect(slide, 0, 4.0, SLIDE_W, 0.08, C.teal)

    # Eyebrow
    add_text_box(
        slide, 0.6, 4.25, 6, 0.30,
        [Run(text="WEEKLY PROGRESS REPORT", size=11, bold=True, color=C.teal, char_spacing=8)],
        valign="middle", margin=0,
    )
    # Big WEEK NN
    add_text_box(
        slide, 0.6, 4.55, 6, 1.0,
        [Run(text=f"WEEK {state.week_no_str}", size=56, bold=True, color=C.white, char_spacing=2)],
        valign="middle", margin=0,
    )
    # Subtitle
    add_text_box(
        slide, 0.6, 5.55, 8, 0.32,
        [Run(text=state.project.short_name, size=14, color="C8D4E3")],
        valign="middle", margin=0,
    )
    # Location (uses the cover-only short location string)
    add_text_box(
        slide, 0.6, 5.88, 8, 0.28,
        [Run(text=f"Plot No. {state.project.plot}   |   {state.project.location_cover}",
             size=10, color="8FA8C2")],
        valign="middle", margin=0,
    )

    # Right-side meta card
    card_x, card_y, card_w = 8.6, 4.45, 4.2
    add_rect(slide, card_x, card_y, card_w, 2.55, C.navy)
    add_rect(slide, card_x, card_y, 0.06, 2.55, C.teal)

    meta_rows = [
        ("REPORTING PERIOD",  state.period),
        ("ISSUE DATE",        state.issue_date),
        ("REPORT REFERENCE",  state.report_ref),
        ("CONTRACT VALUE",    state.project.contract_value_orig),
        ("PLANNED DURATION",  f"{state.project.commencement_full} → {state.project.planned_completion_full}"),
    ]
    for i, (label, value) in enumerate(meta_rows):
        ry = card_y + 0.18 + i * 0.47
        add_text_box(
            slide, card_x + 0.28, ry, card_w - 0.4, 0.18,
            [Run(text=label, size=8, bold=True, color=C.teal, char_spacing=3)],
            valign="middle", margin=0,
        )
        add_text_box(
            slide, card_x + 0.28, ry + 0.18, card_w - 0.4, 0.22,
            [Run(text=value, size=11, bold=True, color=C.white)],
            valign="middle", margin=0,
        )

    # Bottom strip: parties
    strip_y = 7.10
    add_rect(slide, 0, strip_y, SLIDE_W, 0.40, C.cover_strip)
    add_text_box(
        slide, 0.4, strip_y, SLIDE_W - 0.8, 0.40,
        [
            Run(text="ISSUED BY ", bold=True, color=C.teal, char_spacing=2, size=9),
            Run(text=state.project.contractor_short, color=C.white, size=9),
            Run(text="      EMPLOYER ", bold=True, color=C.teal, char_spacing=2, size=9),
            Run(text=state.project.employer_short, color=C.white, size=9),
            Run(text="      CONSULTANT ", bold=True, color=C.teal, char_spacing=2, size=9),
            Run(text=state.project.consultant, color=C.white, size=9),
        ],
        valign="middle", margin=0,
    )

    # Logos top-right (over the photo, on a faint white plate)
    add_rect(slide, 9.55, 0.20, 3.55, 0.85, C.white, line_hex=C.white, line_width=0.5)
    add_image(slide, ASSETS["logo_ew"],  9.65,  0.30, 1.05, 0.65)
    add_image(slide, ASSETS["logo_ah"],  10.85, 0.36, 1.05, 0.50)
    add_image(slide, ASSETS["logo_epc"], 12.10, 0.30, 0.90, 0.65)

    return slide


# ---------- SLIDE 2 — PROJECT OVERVIEW ---------------------------------------

def build_slide_overview(prs, state: WPRState):
    slide = add_blank_slide(prs)
    set_slide_background(slide, C.page)
    add_header(
        slide,
        title="PROJECT OVERVIEW",
        subtitle=f"{state.project.short_name_plain}   |   Architectural Rendering",
    )
    add_footer(slide, 1, period=state.period_short,
               company=state.project.contractor_short,
               report_no=f"Weekly Report No. {state.week_no_str.zfill(3)}")

    # Building rendering left
    add_image(slide, ASSETS["building"], 0.45, 1.15, 8.4, 4.5)
    add_text_box(
        slide, 0.45, 5.70, 8.4, 0.30,
        [Run(text="Architectural rendering — proposed elevation viewed from waterfront approach.",
             size=9.5, italic=True, color=C.slate)],
        align="center", valign="middle", margin=0,
    )

    # Right: project facts panel
    px, pw = 9.10, 3.85
    add_rect(slide, px, 1.15, pw, 5.85, C.panel, line_hex=C.line, line_width=0.75)
    add_rect(slide, px, 1.15, pw, 0.50, C.navy)
    add_text_box(
        slide, px + 0.18, 1.15, pw - 0.3, 0.50,
        [Run(text="PROJECT AT A GLANCE", size=10.5, bold=True, color=C.white, char_spacing=4)],
        valign="middle", margin=0,
    )

    facts = [
        ("BUILDING TYPE",   "B + G + 4 Typical Floor + Roof"),
        ("USE",             "Residential Building"),
        ("LOCATION",        state.project.location),
        ("PLOT NUMBER",     state.project.plot),
        ("EMPLOYER",        state.project.employer_overview),
        ("CONSULTANT",      state.project.consultant),
        ("MAIN CONTRACTOR", "Emirates Pearl Construction\nCCS Company"),
        ("CONTRACT VALUE",  state.project.contract_value_orig),
        ("DURATION",        state.project.duration_full),
    ]
    fy = 1.78
    for label, value in facts:
        add_text_box(
            slide, px + 0.20, fy, pw - 0.4, 0.18,
            [Run(text=label, size=8, bold=True, color=C.teal, char_spacing=2.5)],
            valign="middle", margin=0,
        )
        lines = value.count("\n") + 1
        val_h = 0.22 if lines == 1 else 0.40
        add_text_box(
            slide, px + 0.20, fy + 0.18, pw - 0.4, val_h,
            [Run(text=value, size=10, bold=True, color=C.ink)],
            valign="top", margin=0,
        )
        fy += 0.18 + val_h + 0.10

    return slide


# ---------- SLIDE 3 — PROJECT INFORMATION SUMMARY ----------------------------

def build_slide_info_summary(prs, state: WPRState):
    slide = add_blank_slide(prs)
    set_slide_background(slide, C.page)
    add_header(
        slide,
        title="PROJECT INFORMATION SUMMARY",
        subtitle=f"Week {state.week_no_str}  |  Reporting period {state.period_short}",
    )
    add_footer(slide, 2, period=state.period_short,
               company=state.project.contractor_short,
               report_no=f"Weekly Report No. {state.week_no_str.zfill(3)}")

    # ---- LEFT: parties + contract value pair ----
    section_label(slide, 0.45, 1.10, 5.5, "PROJECT PARTIES")
    parties = [
        ("LOCATION",              state.project.location_inline),
        ("EMPLOYER",              state.project.employer),
        ("CONSULTANT / ENGINEER", state.project.consultant),
        ("MAIN CONTRACTOR",       state.project.contractor),
    ]
    py = 1.45
    for label, value in parties:
        lines = value.count("\n") + 1
        card_h = 0.55 if lines == 1 else 0.75
        add_rect(slide, 0.45, py, 5.7, card_h, C.panel, line_hex=C.line, line_width=0.5)
        add_rect(slide, 0.45, py, 0.06, card_h, C.teal)
        add_text_box(
            slide, 0.65, py + 0.06, 5.4, 0.18,
            [Run(text=label, size=8, bold=True, color=C.teal, char_spacing=2.5)],
            valign="middle", margin=0,
        )
        add_text_box(
            slide, 0.65, py + 0.22, 5.4, card_h - 0.26,
            [Run(text=value, size=11, bold=True, color=C.ink)],
            valign="top", margin=0,
        )
        py += card_h + 0.10

    # Contract value pair
    cv_y = py + 0.05
    for i, label in enumerate(["Contract Value (Original)", "Contract Value (Revised)"]):
        cx = 0.45 + i * 2.85
        add_rect(slide, cx, cv_y, 2.75, 0.65, C.navy)
        add_text_box(
            slide, cx + 0.15, cv_y + 0.06, 2.55, 0.20,
            [Run(text=label.upper(), size=7.5, bold=True, color=C.teal, char_spacing=2)],
            valign="middle", margin=0,
        )
        value = state.project.contract_value_orig if i == 0 else state.project.contract_value_revised
        add_text_box(
            slide, cx + 0.15, cv_y + 0.24, 2.55, 0.32,
            [Run(text=value, size=14, bold=True, color=C.white)],
            valign="middle", margin=0,
        )

    # ---- RIGHT: KPIs + sub-contractors ----
    section_label(slide, 6.40, 1.10, 6.5, "PROGRAMME AT A GLANCE")

    kpis = [
        (state.project.commencement,         "Commencement",        C.steel, 16),
        (state.project.planned_completion,   "Planned Completion",  C.teal, 16),
        (state.project.duration,             "Contract Duration",   C.steel, 16),
        (f"Week {state.week_no_str}",        "Current Period",      C.teal, 18),
        (state.time_elapsed_pct,             "Time Elapsed",        C.steel, 22),
        (state.project.contract_value_short, "Contract Value",      C.teal, 18),
    ]
    col_w, col_gap, base_x, base_y = 2.10, 0.10, 6.40, 1.45
    for i, (v, l, fill, vsize) in enumerate(kpis):
        cx = base_x + (i % 3) * (col_w + col_gap)
        cy = base_y + (i // 3) * 1.05
        _add_kpi(slide, cx, cy, col_w, 0.95, v, l, fill, value_size=vsize)

    # ---- Sub-contractors table ----
    section_label(slide, 6.40, 3.75, 6.5, "APPOINTED SUB-CONTRACTORS")
    _add_subs_table(slide, state, 6.40, 4.10, 6.55)

    return slide


def _add_subs_table(slide, state: WPRState, x: float, y: float, w: float):
    rows = len(state.subs) + 1
    cols = 3
    col_widths = [2.05, 3.20, 1.30]
    row_h = 0.28
    table_h = row_h * rows

    table_shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(table_h))
    table = table_shape.table
    for ci, cw in enumerate(col_widths):
        table.columns[ci].width = Inches(cw)
    for ri in range(rows):
        table.rows[ri].height = Inches(row_h)

    headers = ["TRADE / SCOPE", "SUB-CONTRACTOR", "STATUS"]
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        fill_cell(cell, C.navy)
        align = "center" if ci == 2 else "left"
        style_cell_text(cell, h, color=C.white, bold=True, size=9.5,
                        align=align, valign="middle", char_spacing=2)
        set_cell_border(cell, C.line, 0.5)

    for ri, sub in enumerate(state.subs, start=1):
        stripe = C.white if ri % 2 == 1 else C.panel
        for ci, val in enumerate([sub.trade, sub.name, "● " + sub.status]):
            cell = table.cell(ri, ci)
            fill_cell(cell, stripe)
            if ci == 0:
                style_cell_text(cell, val, color=C.ink, bold=True, size=9.5, align="left", valign="middle")
            elif ci == 1:
                style_cell_text(cell, val, color=C.graphite, size=9.5, align="left", valign="middle")
            else:
                style_cell_text(cell, val, color=C.success, bold=True, size=9.5, align="center", valign="middle")
            set_cell_border(cell, C.line, 0.5)
